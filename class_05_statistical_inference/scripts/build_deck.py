"""Builds the Class 05 instructor deck: ../slides/Class05_Statistical_Inference.pptx

Every number on the slides is computed here from the Class 03 campus survey
(the 518 recorded exam scores are treated as "the whole population"), so the deck
and the student activity notebook always agree.  Edit this file, not the .pptx.

Needs: python-pptx, numpy, pandas, scipy.
"""
import os

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "..", "class_03_eda_deep_dive", "datasets", "campus_survey_clean.csv")
OUT = os.path.join(HERE, "..", "slides", "Class05_Statistical_Inference.pptx")

# ------------------------------------------------------------------ numbers
df = pd.read_csv(DATA)
pop = df["exam_score"].dropna().to_numpy()
MU, SIGMA, N_POP = pop.mean(), pop.std(ddof=0), len(pop)

rng = np.random.default_rng(2026)
means30 = np.array([rng.choice(pop, 30).mean() for _ in range(5000)])
SD_MEANS30 = means30.std()
edges = np.arange(56.5, 73.6, 1.0)
counts, _ = np.histogram(means30, bins=edges)
pct = counts / counts.sum() * 100
centers = [str(int(c)) for c in (edges[:-1] + edges[1:]) / 2]

SE_N = [5, 10, 30, 100, 400]
SE_V = [SIGMA / np.sqrt(n) for n in SE_N]

# 20 confidence intervals (n = 30), choose a seed that gives exactly one miss
def make_cis(seed):
    r = np.random.default_rng(seed)
    out = []
    for _ in range(20):
        s = r.choice(pop, 30, replace=False)
        se = s.std(ddof=1) / np.sqrt(30)
        t = stats.t.ppf(0.975, 29)
        out.append((s.mean(), s.mean() - t * se, s.mean() + t * se))
    return out
for seed in range(1, 500):
    CIS = make_cis(seed)
    if sum(not (lo <= MU <= hi) for _, lo, hi in CIS) == 1:
        break
N_MISS = sum(not (lo <= MU <= hi) for _, lo, hi in CIS)

# worked CI example (n = 40, one random sample)
r40 = np.random.default_rng(11)
s40 = r40.choice(pop, 40, replace=False)
X40, S40 = s40.mean(), s40.std(ddof=1)
SE40 = S40 / np.sqrt(40)
T40 = stats.t.ppf(0.975, 39)
LO40, HI40 = X40 - T40 * SE40, X40 + T40 * SE40

# hypothesis test: heavy job (> 15 h) vs no job
job = df[df.part_time_job_hours_week > 15].exam_score.dropna().to_numpy()
nojob = df[df.part_time_job_hours_week == 0].exam_score.dropna().to_numpy()
tt = stats.ttest_ind(nojob, job, equal_var=False)
sp = np.sqrt(((len(job) - 1) * job.var(ddof=1) + (len(nojob) - 1) * nojob.var(ddof=1)) / (len(job) + len(nojob) - 2))
D_JOB = (nojob.mean() - job.mean()) / sp
DIFF = nojob.mean() - job.mean()
se_diff = np.sqrt(job.var(ddof=1) / len(job) + nojob.var(ddof=1) / len(nojob))
CI_DIFF = (DIFF - 1.96 * se_diff, DIFF + 1.96 * se_diff)

hostel = df[df.residence == "Hostel"].exam_score.dropna().to_numpy()
home = df[df.residence == "Home"].exam_score.dropna().to_numpy()
p_res = stats.ttest_ind(hostel, home, equal_var=False).pvalue
d_res = (hostel.mean() - home.mean()) / np.sqrt((hostel.var(ddof=1) + home.var(ddof=1)) / 2)

# same effect (3 points, SD 10 -> d = 0.30), different sample sizes
SAME_D = []
for n in (20, 100, 1000):
    p = stats.ttest_ind_from_stats(63, 10, n, 60, 10, n).pvalue
    SAME_D.append((n, p))


def fmt_p(p):
    if p < 0.0001:
        return "< 0.0001"
    return f"{p:.3f}" if p >= 0.001 else f"{p:.4f}"


# ------------------------------------------------------------------ style
DARK, LIGHT, CARD = "0F3D3E", "F6F8F7", "FFFFFF"
INK, MUTED = "17302F", "5C6F6E"
TEAL, MINT, AMBER, RED = "1F7A78", "CFE5E1", "F2A541", "C8453F"
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide(dark=False, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK if dark else LIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def rect(s, x, y, w, h, fill=CARD, line=None, rounded=True, shadow=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if not shadow:
        shp.shadow.inherit = False
    return shp


def circle(s, x, y, d, fill):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    return c


def text(s, x, y, w, h, paras, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=6):
    """paras: str | list of str | list of dict(text|runs, size, bold, color, font, italic, space)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("space", space))
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", p.get("font", font))
            f.size = Pt(o.get("size", p.get("size", size)))
            f.bold = o.get("bold", p.get("bold", bold))
            f.italic = o.get("italic", p.get("italic", False))
            f.color.rgb = rgb(o.get("color", p.get("color", color)))
    return tb


def title(s, tag, head, dark=False):
    text(s, 0.6, 0.42, 12, 0.3, tag.upper(), size=12, color=AMBER if dark else TEAL, bold=True)
    text(s, 0.6, 0.72, 12.1, 0.8, head, size=34, color="FFFFFF" if dark else DARK, font=HEAD, bold=True)


def card(s, x, y, w, h, head, body, head_color=TEAL, size=15, fill=CARD, line=MINT):
    rect(s, x, y, w, h, fill=fill, line=line)
    text(s, x + 0.25, y + 0.2, w - 0.5, 0.4, head, size=17, color=head_color, bold=True, font=HEAD)
    text(s, x + 0.25, y + 0.68, w - 0.5, h - 0.85, body, size=size, color=INK)


def style_chart(ch, cat_title, val_title):
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = BODY
    ch.font.size = Pt(12)
    ch.font.color.rgb = rgb(MUTED)
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = rgb(MINT)
    va.format.line.fill.background()
    ca.format.line.color.rgb = rgb(MUTED)
    for ax, t in ((ca, cat_title), (va, val_title)):
        ax.has_title = True
        ax.axis_title.text_frame.text = t
        r = ax.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(12), False, BODY
        r.font.color.rgb = rgb(MUTED)


# ================================================================== 1 TITLE
s = new_slide(dark=True, notes=(
    "TIMING GUIDE (~60-75 min if you do everything; each section stands alone so you can stop anywhere):\n"
    "Slides 2-4 population/sample (8 min) | 5-6 sampling distribution (12 min) | 7-8 confidence intervals (12 min) | "
    "9-12 hypothesis tests & p-values (15 min) | 13-15 effect sizes & mistakes (10 min) | 16 group activity.\n\n"
    "IF YOU ONLY HAVE ~20 MIN left after the presentations: do slides 2, 5, 7 and 9, then hand out the activity.\n\n"
    "All numbers use the Class 03 campus survey. Tell them: 'Pretend these 518 exam scores are every student in the university.' "
    "That way we KNOW the truth and can check our inferences."))
text(s, 0.8, 1.6, 7, 0.4, "DATA SCIENCE · CLASS 05", size=14, color=AMBER, bold=True)
text(s, 0.8, 2.1, 7.2, 2.2, "From a sample\nto the truth", size=54, color="FFFFFF", font=HEAD, bold=True)
text(s, 0.8, 4.5, 6.8, 1.4,
     "Statistical inference: samples, sampling distributions, confidence intervals, hypothesis tests, p-values and effect sizes",
     size=18, color=MINT)
r5 = np.random.default_rng(3)
picked = set(r5.choice(14 * 9, 22, replace=False).tolist())
for i in range(14 * 9):
    cx, cy = i % 14, i // 14
    circle(s, 8.2 + cx * 0.33, 1.5 + cy * 0.5, 0.2, AMBER if i in picked else TEAL)

# ================================================================== 2 QUESTION
s = new_slide(notes=(
    "SAY: 'Last class: probability, where we KNEW the machine (a fair die, a 1% disease) and asked what data to expect. "
    "Inference runs backwards: we have the data and want to learn about the machine.'\n"
    "ASK: 'What is the average exam score of every student at this university? How could we find out without testing everyone?'\n"
    "Take 2-3 answers (survey, sample, ask the registrar). Land on: we look at a few and have to say how sure we are."))
title(s, "The big idea", "We study a few to learn about everyone")
card(s, 0.6, 1.9, 3.8, 3.0, "1 · The question",
     "What is the average exam score of ALL students?", size=22)
card(s, 4.77, 1.9, 3.8, 3.0, "2 · The data",
     "We only have 40 scores.\nWe cannot ask everyone.", size=22)
card(s, 8.93, 1.9, 3.8, 3.0, "3 · The honest answer",
     "An estimate, plus how wrong it could be.", size=22, head_color=RED)
rect(s, 0.6, 5.3, 12.13, 1.3, fill=DARK, rounded=True)
text(s, 0.9, 5.3, 11.6, 1.3, "Inference = using a sample to say something about a population, and being honest about the uncertainty.",
     size=20, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, font=HEAD, bold=True)

# ================================================================== 3 POP VS SAMPLE
s = new_slide(notes=(
    f"Population in our example: {N_POP} exam scores, mean mu = {MU:.1f}, SD sigma = {SIGMA:.1f}. This is the 'truth'.\n"
    f"A sample of 40 gave x-bar = {X40:.1f}, s = {S40:.1f}. Close to the truth but not equal.\n"
    "KEY LINES: parameter = fixed number about the population (usually unknown); statistic = number computed from the sample (changes every time).\n"
    "Greek letters for the population, Latin letters for the sample is a handy memory trick.\n"
    "ASK: 'If I take another sample of 40, will x-bar be the same?' (No: it moves. How much? That is the next slide.)"))
title(s, "1 · Population vs sample", "Parameter (the truth) vs statistic (our estimate)")
for x, head, sub, rows, col in [
    (0.6, "Population", "Everyone we care about", [("Mean", "μ (mu)"), ("Std deviation", "σ (sigma)"), ("Proportion", "p")], TEAL),
    (7.0, "Sample", "The few we actually measured", [("Mean", "x̄ (x-bar)"), ("Std deviation", "s"), ("Proportion", "p̂ (p-hat)")], AMBER),
]:
    rect(s, x, 1.8, 5.73, 3.4, fill=CARD, line=MINT)
    text(s, x + 0.3, 2.0, 5.1, 0.5, head, size=26, color=col if col == TEAL else "B36B00", font=HEAD, bold=True)
    text(s, x + 0.3, 2.55, 5.1, 0.4, sub, size=15, color=MUTED)
    for i, (k, v) in enumerate(rows):
        text(s, x + 0.3, 3.2 + i * 0.55, 2.6, 0.4, k, size=18, color=INK)
        text(s, x + 3.0, 3.2 + i * 0.55, 2.5, 0.4, v, size=20, color=DARK, bold=True)
text(s, 6.05, 3.0, 0.9, 0.6, "→", size=40, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.5, 5.73, 1.2, [
    {"runs": [("Parameter: ", {"bold": True, "color": TEAL}), ("a fixed number, usually unknown.", {})]},
    {"text": f"Our class data: μ = {MU:.1f}"}], size=17)
text(s, 7.0, 5.5, 5.73, 1.2, [
    {"runs": [("Statistic: ", {"bold": True, "color": "B36B00"}), ("computed from data, changes sample to sample.", {})]},
    {"text": f"One sample of 40: x̄ = {X40:.1f}"}], size=17)

# ================================================================== 4 SAMPLING METHOD
s = new_slide(notes=(
    "MESSAGE: how you sample matters more than how many you sample.\n"
    "STORY: 1936 US presidential election. The Literary Digest magazine mailed a huge survey (millions of replies) and confidently predicted the wrong winner. "
    "Their list came from magazine subscribers, car owners and phone directories, which over-represented wealthier voters in the Depression. "
    "A much smaller, properly random poll got it right.\n"
    "ASK: 'Suppose I survey only the students sitting in the front row. Is that random? What could be biased?'\n"
    "Connect back: Class 03's missing sleep data was also non-random (busy students skipped it)."))
title(s, "2 · Sampling", "A big sample is not the same as a good sample")
card(s, 0.6, 1.8, 5.9, 2.6, "Random sample",
     "Every member has an equal chance of being picked.\nErrors average out, so x̄ is aimed at μ.", size=19)
card(s, 6.83, 1.8, 5.9, 2.6, "Biased sample", [
    "Some members are more likely to be picked.",
    "Volunteers only · friends only · survivors only · the front row",
    "More data does NOT fix it: you just get a precise wrong answer."], head_color=RED, size=17)
rect(s, 0.6, 4.75, 12.13, 1.9, fill=DARK)
text(s, 0.95, 4.95, 11.4, 0.4, "REAL EXAMPLE · US election, 1936", size=13, color=AMBER, bold=True)
text(s, 0.95, 5.4, 11.4, 1.1,
     "A magazine poll with millions of replies predicted the wrong winner. A far smaller random poll got it right. "
     "The magazine’s mailing list leaned wealthy.", size=20, color="FFFFFF")

# ================================================================== 5 SAMPLING DISTRIBUTION
s = new_slide(notes=(
    f"WHAT WE DID: pretend we repeat the study 5,000 times. Each time draw 30 students at random and write down the sample mean. "
    f"This chart is the histogram of those 5,000 means.\n"
    f"POINT 1: it is centred on the truth ({MU:.1f}): sample means are not systematically too high or low.\n"
    f"POINT 2: it is much narrower than the population itself (SD of the means = {SD_MEANS30:.1f} vs population SD {SIGMA:.1f}).\n"
    "POINT 3: it is bell-shaped: that is the Central Limit Theorem from last class.\n"
    "ASK: 'In real life do we get to draw 5,000 samples?' (No! We get one. The theory tells us what the other 4,999 would look like.)"))
title(s, "3 · Sampling distribution", "If we repeated the study 5,000 times")
cd = CategoryChartData()
cd.categories = centers
cd.add_series("% of sample means", [round(v, 2) for v in pct])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.7), Inches(7.6), Inches(5.0), cd)
ch = gf.chart
style_chart(ch, "sample mean of 30 exam scores", "% of the 5,000 samples")
pl = ch.plots[0]
pl.gap_width = 12
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = rgb(TEAL)
text(s, 8.6, 1.8, 4.13, 0.4, "Each bar = how often a sample of 30 gave that mean", size=15, color=MUTED)
for i, (big, small, col) in enumerate([
    (f"{MU:.1f}", "true population mean μ, and the centre of the bars", TEAL),
    (f"{SIGMA:.1f}", "spread of individual students (σ)", MUTED),
    (f"{SD_MEANS30:.1f}", "spread of sample means. Averages are much steadier", "B36B00"),
]):
    y = 2.5 + i * 1.4
    rect(s, 8.6, y, 4.13, 1.2, fill=CARD, line=MINT)
    text(s, 8.8, y, 1.4, 1.2, big, size=32, color=col, font=HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.2, y, 2.4, 1.2, small, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================== 6 STANDARD ERROR
s = new_slide(notes=(
    "DEFINITION: the standard error (SE) is the standard deviation of the sampling distribution: how far a sample mean typically lands from the truth.\n"
    f"FORMULA: SE = sigma / sqrt(n). With sigma = {SIGMA:.1f}: n=5 gives {SE_V[0]:.1f}, n=30 gives {SE_V[2]:.1f}, n=100 gives {SE_V[3]:.1f}.\n"
    "KEY CONSEQUENCE: diminishing returns. Four times as much data only halves the error. Going from 100 to 400 students costs a lot more than 5 to 20.\n"
    "WARNING: SD describes the spread of individual data; SE describes the precision of an estimate. Students mix these up constantly.\n"
    "ASK: 'To cut our error in a third, how much more data do we need?' (9 times.)"))
title(s, "3 · Sampling distribution", "Standard error: how wrong is a typical x̄?")
rect(s, 0.6, 1.8, 5.2, 2.2, fill=DARK)
text(s, 0.6, 1.8, 5.2, 2.2, "SE = σ / √n", size=44, color="FFFFFF", font=HEAD, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.3, 5.2, 2.4, [
    {"runs": [("SD", {"bold": True, "color": TEAL}), (" = spread of individual students.", {})]},
    {"runs": [("SE", {"bold": True, "color": "B36B00"}), (" = how precise our estimate of the mean is.", {})]},
    {"runs": [("4× the data", {"bold": True}), (" → half the error.", {})]},
], size=18, space=10)
cd = CategoryChartData()
cd.categories = [f"n = {n}" for n in SE_N]
cd.add_series("Standard error", [round(v, 2) for v in SE_V])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.2), Inches(1.7), Inches(6.5), Inches(5.0), cd)
ch = gf.chart
style_chart(ch, "sample size", f"standard error of the mean (σ = {SIGMA:.1f})")
pl = ch.plots[0]
pl.gap_width = 60
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = rgb(TEAL)
pl.has_data_labels = True
pl.data_labels.number_format = "0.0"
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(14)
pl.data_labels.font.bold = True
pl.data_labels.font.color.rgb = rgb(DARK)

# ================================================================== 7 CI PICTURE
s = new_slide(notes=(
    f"WHAT YOU SEE: 20 different random samples of 30 students. For each we computed a 95% confidence interval. The dashed line is the true mean ({MU:.1f}), which we know only because we are pretending the 518 scores are everyone.\n"
    f"{20 - N_MISS} of the 20 intervals catch the truth, {N_MISS} (red) missed. That is what '95%' means: the METHOD works about 95% of the time in the long run.\n"
    "LANGUAGE TRAP: after you compute ONE interval, it either contains the truth or it does not. We do not say 'there is a 95% chance mu is in here'. We say 'we used a method that captures the truth 95% of the time'.\n"
    "ASK: 'Which would you prefer: a 99% interval or a 95% one? What is the price?' (Wider, less precise.)"))
title(s, "4 · Confidence intervals", "A net, not a bullseye: 20 samples, 20 intervals")
X0, SC = 0.9, 7.4 / 18.0
def vx(v):
    return X0 + (min(max(v, 56), 74) - 56) * SC
for i, (m, lo, hi) in enumerate(CIS):
    y = 1.95 + i * 0.225
    ok = lo <= MU <= hi
    col = TEAL if ok else RED
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(vx(lo)), Inches(y), Inches(vx(hi)), Inches(y))
    ln.line.color.rgb = rgb(col)
    ln.line.width = Pt(3)
    circle(s, vx(m) - 0.06, y - 0.06, 0.12, col)
mu_line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(vx(MU)), Inches(1.78), Inches(vx(MU)), Inches(6.4))
mu_line.line.color.rgb = rgb(AMBER)
mu_line.line.width = Pt(2.5)
mu_line.line.dash_style = 4  # dash
for v in (56, 60, 64, 68, 72):
    text(s, vx(v) - 0.3, 6.45, 0.6, 0.3, str(v), size=12, color=MUTED, align=PP_ALIGN.CENTER)
text(s, vx(MU) - 1.0, 1.5, 2.0, 0.3, f"true μ = {MU:.1f}", size=13, color="B36B00", bold=True, align=PP_ALIGN.CENTER)
rect(s, 8.9, 1.8, 3.83, 1.5, fill=CARD, line=MINT)
text(s, 9.05, 1.8, 1.7, 1.5, f"{20 - N_MISS}/20", size=32, color=TEAL, font=HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 10.8, 1.8, 1.8, 1.5, "intervals caught the true mean", size=15, anchor=MSO_ANCHOR.MIDDLE)
card(s, 8.9, 3.55, 3.83, 1.45, "95% is about the method",
     "Repeat the process many times: about 95% of intervals catch μ.", size=14)
card(s, 8.9, 5.2, 3.83, 1.45, "Price of confidence",
     "99% → wider interval. More data → narrower.", size=14)

# ================================================================== 8 CI RECIPE
s = new_slide(notes=(
    f"WORKED EXAMPLE: one random sample of 40 scores: x-bar = {X40:.1f}, s = {S40:.1f}.\n"
    f"SE = s/sqrt(40) = {SE40:.2f}.  t* (39 df, 95%) = {T40:.2f}.  Margin of error = {T40 * SE40:.2f}.  Interval = [{LO40:.1f}, {HI40:.1f}].\n"
    f"CHECK: true mean {MU:.1f} is {'inside' if LO40 <= MU <= HI40 else 'outside'} the interval.\n"
    "WHY t AND NOT z: we had to estimate sigma with s, which adds uncertainty, so t has slightly fatter tails. For n > 30 it is almost the same as 1.96.\n"
    "ASK: 'If I take 4x as many students, what happens to the margin of error?' (Halves.)\n"
    "Reading aloud the correct sentence is worth practising: 'We are 95% confident that the true mean is between ... and ...'."))
title(s, "4 · Confidence intervals", "How to build one (95%)")
rect(s, 0.6, 1.8, 6.2, 1.5, fill=DARK)
text(s, 0.6, 1.8, 6.2, 1.5, "x̄  ±  t* × s / √n", size=36, color="FFFFFF", font=HEAD, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
steps = [
    ("Sample mean", f"x̄ = {X40:.1f}  (n = 40)"),
    ("Spread of the sample", f"s = {S40:.1f}"),
    ("Standard error", f"s / √n = {SE40:.2f}"),
    ("Critical value", f"t* = {T40:.2f}  (95%, 39 df)"),
    ("Margin of error", f"{T40:.2f} × {SE40:.2f} = ±{T40 * SE40:.1f}"),
]
for i, (k, v) in enumerate(steps):
    y = 3.55 + i * 0.62
    circle(s, 0.6, y + 0.02, 0.42, AMBER)
    text(s, 0.6, y + 0.02, 0.42, 0.42, str(i + 1), size=15, bold=True, color=DARK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)
    text(s, 1.25, y, 2.5, 0.5, k, size=16, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.7, y, 3.2, 0.5, v, size=17, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 7.2, 1.8, 5.53, 2.3, fill=CARD, line=TEAL)
text(s, 7.5, 1.95, 5.0, 0.4, "THE RESULT", size=13, color=TEAL, bold=True)
text(s, 7.5, 2.4, 5.0, 1.2, f"{LO40:.1f}  to  {HI40:.1f}", size=40, color=DARK, font=HEAD, bold=True)
text(s, 7.5, 3.4, 5.0, 0.6, f"(true μ = {MU:.1f}, inside)" if LO40 <= MU <= HI40 else f"(true μ = {MU:.1f}, missed this time)",
     size=15, color=MUTED)
card(s, 7.2, 4.4, 2.65, 2.3, "Say", "“We are 95% confident the mean is between these two numbers.”", size=14)
card(s, 10.08, 4.4, 2.65, 2.3, "Don’t say", "“95% of students score between these.” That is a different thing.", head_color=RED, size=14)

# ================================================================== 9 HYPOTHESIS TEST FRAMEWORK
s = new_slide(notes=(
    "COURTROOM ANALOGY: H0 is 'innocent until proven guilty'. We assume nothing is going on, then ask whether the evidence is so surprising under that assumption that we reject it.\n"
    "'Fail to reject' is NOT 'proved H0 true', exactly like 'not guilty' is not 'innocent'. Insufficient evidence is not evidence of absence.\n"
    "THE 5 STEPS are the template for every test they will ever do (t-test, chi-square, anova, ...).\n"
    "ASK: 'Claim: students with heavy jobs score lower. What is H0? What is H1?' H0: the mean scores are equal. H1: they differ."))
title(s, "5 · Hypothesis testing", "Innocent until proven guilty")
card(s, 0.6, 1.8, 5.4, 2.5, "The courtroom", [
    "H₀ (null): nothing is going on.",
    "H₁ (alternative): something is going on.",
    "Data = the evidence.",
], size=16)
card(s, 0.6, 4.55, 5.4, 2.15, "The verdict", [
    "Evidence very surprising under H₀ → reject H₀.",
    "Otherwise: fail to reject. That is “not guilty”, not “innocent”."], size=15, head_color=RED)
for i, (k, v) in enumerate([
    ("State the hypotheses", "H₀: μ₁ = μ₂   vs   H₁: μ₁ ≠ μ₂"),
    ("Pick a threshold", "α = 0.05 (decide BEFORE looking)"),
    ("Compute a test statistic", "How many SEs apart are the groups?"),
    ("Find the p-value", "How surprising is that if H₀ is true?"),
    ("Decide and report", "Reject / fail to reject, plus estimate and size"),
]):
    y = 1.8 + i * 0.98
    rect(s, 6.4, y, 6.33, 0.85, fill=CARD, line=MINT)
    circle(s, 6.6, y + 0.2, 0.45, TEAL)
    text(s, 6.6, y + 0.2, 0.45, 0.45, str(i + 1), size=16, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)
    text(s, 7.3, y + 0.08, 5.3, 0.4, k, size=17, bold=True, color=DARK)
    text(s, 7.3, y + 0.43, 5.3, 0.4, v, size=14, color=MUTED)

# ================================================================== 10 TEST IN ACTION
s = new_slide(notes=(
    f"QUESTION: do students who work more than 15 hours a week score differently from students with no job?\n"
    f"DATA: no job n = {len(nojob)}, mean {nojob.mean():.1f}; heavy job n = {len(job)}, mean {job.mean():.1f}. Difference = {DIFF:.1f} points.\n"
    f"Welch two-sample t-test: t = {tt.statistic:.1f}, p = {tt.pvalue:.1e}. 95% CI for the difference: [{CI_DIFF[0]:.1f}, {CI_DIFF[1]:.1f}].\n"
    "INTERPRET: if heavy jobs made NO difference, a gap this big would almost never come from sampling luck alone. So we reject H0.\n"
    "CAUTION: this is observational: heavy workers also sleep less (Class 03). A small p does not say jobs CAUSE lower scores.\n"
    "Code: scipy.stats.ttest_ind(a, b, equal_var=False)."))
title(s, "5 · Hypothesis testing", "In action: does a heavy job change exam scores?")
for x, head, n_, m_, col in [
    (0.6, "No job", len(nojob), nojob.mean(), TEAL),
    (4.75, "Job > 15 h / week", len(job), job.mean(), RED),
]:
    rect(s, x, 1.8, 3.8, 2.4, fill=CARD, line=MINT)
    text(s, x + 0.25, 1.95, 3.3, 0.4, head, size=17, bold=True, color=DARK, font=HEAD)
    text(s, x + 0.25, 2.5, 3.3, 1.0, f"{m_:.1f}", size=48, color=col, font=HEAD, bold=True)
    text(s, x + 0.25, 3.55, 3.3, 0.4, f"mean score · n = {n_}", size=14, color=MUTED)
rect(s, 8.9, 1.8, 3.83, 2.4, fill=DARK)
text(s, 9.15, 1.95, 3.4, 0.4, "WELCH t-TEST", size=13, color=AMBER, bold=True)
text(s, 9.15, 2.4, 3.4, 1.7, [
    {"runs": [("t = ", {"color": MINT}), (f"{tt.statistic:.1f}", {"bold": True})]},
    {"runs": [("p ", {"color": MINT}), (fmt_p(tt.pvalue), {"bold": True})]},
], size=30, color="FFFFFF", font=HEAD, space=2)
card(s, 0.6, 4.5, 6.0, 2.2, "Reading it", [
    f"Gap = {DIFF:.1f} points  (95% CI {CI_DIFF[0]:.1f} to {CI_DIFF[1]:.1f}).",
    "If jobs made no difference, a gap this large from luck alone would be extremely rare → reject H₀."], size=15)
card(s, 6.83, 4.5, 5.9, 2.2, "Careful", [
    "Significant ≠ caused by the job.",
    "Heavy workers also sleep less. Other differences can hide behind the gap."], head_color=RED, size=15)

# ================================================================== 11 P-VALUE
s = new_slide(dark=True, notes=(
    "THE DEFINITION IN PLAIN WORDS: 'Assume nothing is going on. How often would luck alone produce a result at least this extreme?' That fraction is the p-value.\n"
    "p = 0.03 means: IF H0 were true, about 3% of studies would show a gap this big or bigger.\n"
    "IT IS NOT: the probability that H0 is true; the probability the result is a fluke; the size or importance of the effect; the probability of replicating.\n"
    "Bayes link (last class): P(data | H0) is not P(H0 | data). This is the same P(A|B) vs P(B|A) trap!\n"
    "ASK: 'p = 0.04 in one study and p = 0.06 in another. Are they really different findings?' (No. 0.05 is a convention, not a cliff.)"))
title(s, "6 · p-values", "The p-value: surprise if nothing is going on", dark=True)
rect(s, 0.6, 1.9, 12.13, 1.55, fill=TEAL)
text(s, 0.95, 1.9, 11.4, 1.55,
     "p-value = P( data this extreme or more  |  H₀ is true )", size=28, color="FFFFFF", font=HEAD, bold=True,
     anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.6, 3.8, 5.95, 2.9, fill="145253")
text(s, 0.95, 3.95, 5.3, 0.4, "IT IS", size=14, color=AMBER, bold=True)
text(s, 0.95, 4.4, 5.3, 2.2, [
    "A measure of surprise under the “nothing is going on” assumption.",
    "Smaller p = harder to blame luck."], size=17, color="FFFFFF", space=10)
rect(s, 6.78, 3.8, 5.95, 2.9, fill="145253")
text(s, 7.13, 3.95, 5.3, 0.4, "IT IS NOT", size=14, color="FF9C93", bold=True)
text(s, 7.13, 4.4, 5.3, 2.2, [
    "The probability that H₀ is true.",
    "The size or importance of the effect.",
    "A magic switch at 0.05."], size=17, color="FFFFFF", space=10)

# ================================================================== 12 ERRORS
s = new_slide(notes=(
    "TWO WAYS TO BE WRONG. Type I = false alarm (we claim an effect that is not there). Type II = miss (a real effect we did not detect).\n"
    "alpha = 0.05 means: when H0 is true we accept a 5% chance of a false alarm. Run 20 tests of nothing and one will come out 'significant' by luck: this is why multiple testing / p-hacking is dangerous.\n"
    "POWER = 1 - P(Type II): the chance of detecting a real effect. Bigger samples and bigger effects raise power.\n"
    "Medical-test link: false positives and false negatives from the Bayes lesson are the same idea.\n"
    "ASK: 'Which error is worse for a new medicine? For a spam filter?'"))
title(s, "6 · p-values", "Two ways to be wrong")
text(s, 3.6, 1.7, 4.4, 0.4, "H₀ is actually TRUE", size=15, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 8.2, 1.7, 4.5, 0.4, "H₀ is actually FALSE", size=15, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 2.55, 2.8, 0.8, "We reject H₀\n(“significant”)", size=16, bold=True, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.7, 2.8, 0.8, "We fail to\nreject H₀", size=16, bold=True, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
for x, y, head, body, col, fillc in [
    (3.6, 2.2, "Type I error", "False alarm. Probability = α (e.g. 5%).", RED, "FBE9E7"),
    (8.2, 2.2, "Correct", "Real effect found. Probability = power.", TEAL, "E3F1EE"),
    (3.6, 4.35, "Correct", "Nothing there, nothing claimed.", TEAL, "E3F1EE"),
    (8.2, 4.35, "Type II error", "A miss. Real effect not detected.", "B36B00", "FFF1DB"),
]:
    rect(s, x, y, 4.5 if x > 5 else 4.4, 1.9, fill=fillc)
    text(s, x + 0.25, y + 0.2, 4.0, 0.4, head, size=19, color=col, bold=True, font=HEAD)
    text(s, x + 0.25, y + 0.8, 4.0, 1.0, body, size=15, color=INK)
text(s, 0.6, 6.5, 12.1, 0.5, "Run 20 tests of pure noise and one will look “significant” at α = 0.05. That is why you decide the test before looking.",
     size=15, color=MUTED)

# ================================================================== 13 EFFECT SIZE
rows_txt = "; ".join(f"n={n} per group: p {fmt_p(p)}" for n, p in SAME_D)
s = new_slide(notes=(
    "p-value answers: 'is there an effect?'  Effect size answers: 'how big is it?'  They are different questions.\n"
    "COHEN'S d = (mean1 - mean2) / pooled SD. It is a difference measured in standard deviations, so it is unit-free and comparable. Rough rules of thumb: 0.2 small, 0.5 medium, 0.8 large (context matters!).\n"
    f"TABLE: the SAME effect (3 points, SD 10, d = 0.30) gives {rows_txt}. Sample size alone flips the verdict while the effect did not change.\n"
    f"OUR DATA: heavy job vs no job: d = {D_JOB:.2f} (medium to large). Hostel vs home residence: d = {d_res:.2f}, p = {p_res:.2f}: nothing there.\n"
    "ASK: 'Which would you rather have: a significant result with d = 0.05, or a non-significant one with d = 0.8 from 10 students?' (Discuss: the second needs more data, the first may not matter.)"))
title(s, "7 · Effect sizes", "Significant is not the same as big")
rect(s, 0.6, 1.8, 6.0, 1.55, fill=DARK)
text(s, 0.6, 1.8, 6.0, 1.55, "d = (x̄₁ − x̄₂) / s_pooled", size=30, color="FFFFFF", font=HEAD, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 3.6, 6.0, 0.5, "Difference measured in standard deviations", size=16, color=MUTED)
for i, (dv, lab) in enumerate([("0.2", "small"), ("0.5", "medium"), ("0.8", "large")]):
    x = 0.6 + i * 2.05
    rect(s, x, 4.25, 1.9, 1.15, fill=CARD, line=MINT)
    text(s, x, 4.3, 1.9, 0.65, dv, size=28, color=TEAL, font=HEAD, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 4.95, 1.9, 0.35, lab, size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.7, 6.0, 1.0, [
    {"runs": [("Our data: ", {"bold": True}),
              (f"heavy job vs none d = {D_JOB:.2f}; hostel vs home d = {d_res:.2f} (p = {p_res:.2f}).", {})]}], size=15)
text(s, 7.0, 1.8, 5.7, 0.5, "Same effect (d = 0.30), only the sample size changes", size=17, bold=True, color=DARK, font=HEAD)
tbl = s.shapes.add_table(4, 3, Inches(7.0), Inches(2.45), Inches(5.73), Inches(2.6)).table
for j, hd in enumerate(["n per group", "p-value", "Verdict at α = 0.05"]):
    c = tbl.cell(0, j)
    c.text = hd
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(DARK)
    for p in c.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14), True, rgb("FFFFFF"), BODY
for i, (n, p) in enumerate(SAME_D, start=1):
    vals = [str(n), fmt_p(p), "significant" if p < 0.05 else "not significant"]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = rgb("FFFFFF" if i % 2 else "EEF4F2")
        for pgh in c.text_frame.paragraphs:
            for r in pgh.runs:
                r.font.size, r.font.name = Pt(16), BODY
                r.font.color.rgb = rgb(RED if (j == 2 and p >= 0.05) else INK)
                r.font.bold = j == 2
text(s, 7.0, 5.35, 5.73, 1.3,
     "A huge sample can make a trivial effect “significant”. A small sample can hide a real one. Always report the size.",
     size=16, color=INK)

# ================================================================== 14 REPORT
s = new_slide(notes=(
    "THE HABIT: always give three things: (1) the estimate with its interval, (2) the p-value if you tested, (3) the effect size, in plain words.\n"
    "Read the example sentence aloud and ask them to find each of the three parts.\n"
    "This is exactly what groups will practise in the activity: they must write one such sentence for their topic."))
title(s, "Putting it together", "How to report a finding")
for i, (head, body, col) in enumerate([
    ("Estimate + interval", f"Gap of {DIFF:.1f} points (95% CI {CI_DIFF[0]:.1f} to {CI_DIFF[1]:.1f})", TEAL),
    ("Test", f"Welch t-test: p {fmt_p(tt.pvalue)}", "B36B00"),
    ("Effect size", f"d = {D_JOB:.2f}, a medium-to-large difference", RED),
]):
    x = 0.6 + i * 4.1
    rect(s, x, 1.8, 3.95, 2.3, fill=CARD, line=MINT)
    circle(s, x + 0.25, 2.0, 0.55, col)
    text(s, x + 0.25, 2.0, 0.55, 0.55, str(i + 1), size=18, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)
    text(s, x + 0.25, 2.75, 3.45, 0.4, head, size=18, bold=True, color=DARK, font=HEAD)
    text(s, x + 0.25, 3.2, 3.45, 0.9, body, size=15, color=INK)
rect(s, 0.6, 4.5, 12.13, 2.0, fill=DARK)
text(s, 0.95, 4.65, 11.4, 0.4, "A COMPLETE SENTENCE", size=13, color=AMBER, bold=True)
text(s, 0.95, 5.1, 11.4, 1.5,
     f"“Students with a job over 15 h/week scored {DIFF:.1f} points lower on average "
     f"(95% CI {CI_DIFF[0]:.1f} to {CI_DIFF[1]:.1f}; Welch t-test p {fmt_p(tt.pvalue)}; d = {D_JOB:.2f}).”",
     size=20, color="FFFFFF", font=HEAD)

# ================================================================== 15 MISTAKES
s = new_slide(notes=(
    "Use this as a closing checklist. Pick 2-3 and ask the class for an example of each.\n"
    "Tie back: 'Which of these did we see in the EDA class?' (Simpson's paradox is the confounding one; missing data is the biased-sample one.)"))
title(s, "Wrap-up", "Five mistakes to avoid")
mist = [
    ("“p = 0.03 means a 3% chance H₀ is true.”", "p is about the data if H₀ is true."),
    ("“Not significant, so there is no effect.”", "Maybe too little data. Look at the interval."),
    ("“Significant, so it matters.”", "Check the effect size and the context."),
    ("“A 95% interval contains 95% of the data.”", "It is about the mean, not individuals."),
    ("“Big sample, so no bias.”", "Bad sampling stays bad at any size."),
]
for i, (a, b) in enumerate(mist):
    y = 1.75 + i * 1.0
    rect(s, 0.6, y, 12.13, 0.88, fill=CARD, line=MINT)
    circle(s, 0.8, y + 0.2, 0.48, RED)
    text(s, 0.8, y + 0.2, 0.48, 0.48, "✕", size=16, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)
    text(s, 1.55, y, 5.6, 0.88, a, size=17, bold=True, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.3, y, 5.3, 0.88, b, size=16, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================== 16 ACTIVITY
s = new_slide(dark=True, notes=(
    "HAND OUT: 05_STUDENT_group_activity.ipynb (Colab). Five groups, one topic each, ~25 min prep + 5 min per group.\n"
    "The notebook gives them the data and the task, but NOT the code: they must work it out (AI allowed, with a log).\n"
    "Population = the 518 exam scores, so every group can compare their answer with the known truth.\n"
    "Questions for you to ask after each group are in 05_INSTRUCTOR_group_activity.ipynb."))
title(s, "Your turn", "Five groups, five questions, one notebook", dark=True)
acts = [
    ("1", "Population vs sample", "Does a random sample find μ? What about a convenience sample?"),
    ("2", "Sampling distribution", "Draw many samples at n = 10, 30, 100. What shrinks, and how fast?"),
    ("3", "Confidence intervals", "Build 100 intervals. How many catch the truth at 90%, 95%, 99%?"),
    ("4", "Hypothesis test & p-value", "Do heavy workers score lower? Test it, then check with shuffling."),
    ("5", "Effect size", "Rank comparisons by d and by p. Same p, different d?"),
]
for i, (n, head, body) in enumerate(acts):
    y = 1.75 + i * 1.0
    rect(s, 0.6, y, 12.13, 0.88, fill="145253")
    circle(s, 0.8, y + 0.19, 0.5, AMBER)
    text(s, 0.8, y + 0.19, 0.5, 0.5, n, size=17, bold=True, color=DARK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)
    text(s, 1.6, y, 3.6, 0.88, head, size=18, bold=True, color="FFFFFF", font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.3, y, 7.3, 0.88, body, size=15, color=MINT, anchor=MSO_ANCHOR.MIDDLE)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("Saved", OUT)
print(f"pop mean {MU:.2f} sd {SIGMA:.2f} N {N_POP} | sd of means(30) {SD_MEANS30:.2f}")
print(f"CI misses {N_MISS}/20 | n=40 mean {X40:.2f} s {S40:.2f} CI [{LO40:.2f},{HI40:.2f}]")
print(f"job test t {tt.statistic:.2f} p {tt.pvalue:.2e} diff {DIFF:.2f} d {D_JOB:.2f} CI {CI_DIFF}")
print("same d:", SAME_D, "| hostel/home d", round(d_res, 3), "p", round(p_res, 3))
