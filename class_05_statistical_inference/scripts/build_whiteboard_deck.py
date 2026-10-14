"""Builds the Class 05 whiteboard deck (Venn-style pictures + maths): ../slides/Class05_Whiteboard_Pictures.pptx

Venn/set pictures where sets really fit (population vs sample, bias, error types) and
"overlap / area" pictures where they don't (sampling distribution, p-value, effect size).
Reuses the styling helpers extracted into class_04's deck_helpers.py.
"""
import os
import sys

import numpy as np
from lxml import etree
from pptx.oxml.ns import qn
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "class_04_probability_foundations", "scripts"))
from deck_helpers import *  # noqa: E402,F401,F403

OUT = os.path.join(HERE, "..", "slides", "Class05_Whiteboard_Pictures.pptx")
C_A, C_B = "1F7A78", "E08A00"
rng = np.random.default_rng(55)


def alpha(shape, pct):
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))


def circ(s, cx, cy, d, color, a=30, line_w=4):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(color)
    alpha(c, a)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(line_w)
    c.shadow.inherit = False
    return c


def frame(s, x, y, w, h):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background()
    r.line.color.rgb = rgb(DARK)
    r.line.width = Pt(3)
    r.shadow.inherit = False


def lab(s, x, y, t, size=20, color=DARK, w=2.4, bold=True):
    text(s, x - w / 2, y - 0.3, w, 0.6, t, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)


def panel(s, head, lines, y=1.8, h=4.9):
    rect(s, 7.6, y, 5.13, h, fill=CARD, line=MINT)
    text(s, 7.9, y + 0.2, 4.6, 0.4, head, size=17, bold=True, color=TEAL, font=HEAD)
    paras = []
    for ln in lines:
        if isinstance(ln, tuple):
            paras.append({"text": ln[0], "size": ln[1], "bold": True, "color": ln[2] if len(ln) > 2 else DARK, "font": HEAD})
        else:
            paras.append({"text": ln, "size": 17})
    text(s, 7.9, y + 0.8, 4.6, h - 1.0, paras, space=10)


def dots(s, cx, cy, r, n, color, size=0.13, seed=0):
    g = np.random.default_rng(seed)
    pts = []
    while len(pts) < n:
        x, y = g.uniform(-r, r, 2)
        if x * x + y * y <= (r - size) ** 2:
            pts.append((x, y))
    for x, y in pts:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + x - size / 2), Inches(cy + y - size / 2),
                               Inches(size), Inches(size))
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(color)
        c.line.fill.background()
        c.shadow.inherit = False
    return pts


def curve_chart(s, x, y, w, h, xs, series, names, colors, xt, yt, label_every=4, kind="line"):
    cd = CategoryChartData()
    cd.categories = [f"{v:g}" if i % label_every == 0 else "" for i, v in enumerate(xs)]
    for nm, vals in zip(names, series):
        cd.add_series(nm, [round(float(v), 4) for v in vals])
    ch = s.shapes.add_chart(XL_CHART_TYPE.AREA if kind == "area" else XL_CHART_TYPE.LINE,
                            Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    style_chart(ch, xt or " ", yt or " ")
    for ser, col in zip(ch.plots[0].series, colors):
        if kind == "area":
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(col)
        else:
            ser.format.line.color.rgb = rgb(col)
            ser.format.line.width = Pt(3.5)
            ser.smooth = True
    return ch


# ================================================================ 1 population vs sample
s = new_slide(notes=(
    "DRAW IT: a big circle = population (all 518 students). Scatter many small dots in it. Circle ~25 dots in a different colour: that is the sample (40 students).\n"
    "SAY: 'The big circle is everyone. The few coloured dots are who we actually measured. The average of the big circle is mu: a fixed number we usually can't see. The average of the coloured dots is x-bar: our estimate. Inference is about how far x-bar can be from mu.'\n"
    "ASK: 'If I pick a different 40, do I get the same x-bar?' (No: the coloured dots change, mu doesn't.)"))
title(s, "Whiteboard 1 · Population vs sample", "The sample is a circle inside the circle")
frame(s, 0.6, 1.8, 6.6, 4.9)
circ(s, 3.9, 4.25, 4.5, C_A, a=15)
dots(s, 3.9, 4.25, 2.25, 140, "9CBFBA", seed=1)
dots(s, 3.9, 4.25, 2.25, 22, C_B, size=0.2, seed=2)
text(s, 0.8, 1.9, 3.0, 0.8, ["Population", "all 518 students"], size=16, bold=True, color=C_A, space=0)
text(s, 5.0, 6.05, 2.1, 0.6, ["● Sample", "40 students"], size=16, bold=True, color=C_B, space=0)
panel(s, "In plain words", [
    ("Population  →  μ", 24, C_A),
    "The average of EVERYONE. A fixed number, usually unknown.",
    ("Sample  →  x̄", 24, "B36B00"),
    "The average of the few we measured. Changes with every new sample.",
    ("Inference = how far can x̄ be from μ?", 19),
])

# ================================================================ 2 random vs biased
s = new_slide(notes=(
    "DRAW IT: two circles of the same population. LEFT: sample dots spread everywhere (random). RIGHT: sample dots all packed in one corner, circled as 'the keenest students' (biased).\n"
    "SAY: 'Random: the sample looks like a small copy of the population, so its average lands near the truth. Biased: the sample only comes from one part, so its average belongs to that part, not to everyone. A bigger biased sample is just a more precise wrong answer.'\n"
    "NUMBERS from our data: random samples average 64.9; the 40 students who study most average about 60.0 (they are mostly Engineering, who score lower)."))
title(s, "Whiteboard 2 · Sampling", "Where the dots come from matters")
frame(s, 0.6, 1.8, 6.6, 4.9)
for cx, seed, title_t, col in [(2.35, 3, "Random", C_A), (5.45, 4, "Biased", RED)]:
    circ(s, cx, 4.05, 2.7, C_A, a=12, line_w=3)
    dots(s, cx, 4.05, 1.35, 70, "9CBFBA", size=0.11, seed=seed)
    lab(s, cx, 2.3, title_t, size=21, color=col, w=2.4)
# random sample: spread
dots(s, 2.35, 4.05, 1.35, 14, C_B, size=0.17, seed=7)
# biased sample: packed in a corner
corner = circ(s, 6.0, 3.45, 1.0, RED, a=25, line_w=3)
g = np.random.default_rng(9)
for _ in range(14):
    a_, r_ = g.uniform(0, 6.28), g.uniform(0, 0.38)
    x, y = 6.0 + r_ * np.cos(a_), 3.45 + r_ * np.sin(a_)
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.085), Inches(y - 0.085), Inches(0.17), Inches(0.17))
    d.fill.solid()
    d.fill.fore_color.rgb = rgb(C_B)
    d.line.fill.background()
    d.shadow.inherit = False
text(s, 5.55, 5.5, 1.7, 0.5, "the keenest\nstudents only", size=13, color=RED, align=PP_ALIGN.CENTER)
lab(s, 2.35, 5.75, "x̄ ≈ 64.9  ✓", size=20, color=C_A, w=2.6)
lab(s, 4.4, 6.35, "truth: μ = 64.9", size=15, color=MUTED, w=3.0, bold=False)
panel(s, "In plain words", [
    ("Random sample", 22, C_A),
    "Everyone has an equal chance → x̄ lands near μ.",
    ("Biased sample", 22, RED),
    "Only one part is picked → x̄ ≈ 60.0 for the 40 keenest students. Wrong, however many you take.",
    "Sampling method beats sample size.",
])

# ================================================================ 3 sampling distribution
s = new_slide(notes=(
    "DRAW IT: a wide bell for individual students and a much narrower bell for sample means, both centred on the same mean.\n"
    "SAY: 'Individual students are spread out: some score 40, some 90. But if I take many samples of 30 and plot their AVERAGES, the averages bunch tightly around the true mean. That narrow curve is the sampling distribution. Its width is the standard error, sigma over root n.'\n"
    "NUMBERS: SD of students 12.2; SE for n=30: 12.2/sqrt(30) = 2.2. Four times the data = half the error."))
title(s, "Whiteboard 3 · Sampling distribution", "Averages bunch up. Individuals don’t")
xs = np.arange(30, 100.1, 2.5)
curve_chart(s, 0.6, 1.7, 6.8, 5.0, xs,
            [stats.norm.pdf(xs, 64.9, 12.2), stats.norm.pdf(xs, 64.9, 2.2)],
            ["individual students (SD 12.2)", "averages of 30 (SE 2.2)"], [C_A, C_B],
            "exam score", "density", label_every=4)
text(s, 5.0, 2.0, 2.4, 0.4, "● averages of 30", size=15, bold=True, color=C_B)
text(s, 5.0, 2.4, 2.4, 0.4, "● individual students", size=15, bold=True, color=C_A)
panel(s, "In plain words", [
    ("Same centre, different width", 21, C_A),
    "SD (12.2): how spread out students are.",
    "SE (2.2): how spread out sample AVERAGES are.",
    ("SE = σ / √n", 26, "B36B00"),
    "4× the data → half the error.",
])

# ================================================================ 4 confidence interval
s = new_slide(notes=(
    "DRAW IT: a number line. Mark mu (true mean) as a dashed vertical line. Draw 6 horizontal intervals (nets) with a dot in the middle, most crossing the line, one not.\n"
    "SAY: 'Each net is built from one sample: sample mean plus or minus a margin. The true mean sits still; the nets move. A 95% net-making method catches the truth about 95 times in 100. This particular net either has it or it does not.'\n"
    "RECIPE: x-bar plus or minus t* x s/sqrt(n). Sample n=40: 65.1 plus or minus 3.3 gives 61.9 to 68.4."))
title(s, "Whiteboard 4 · Confidence interval", "Each sample throws a net")
frame(s, 0.6, 1.8, 6.6, 4.9)
mu_x = 3.9
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(mu_x), Inches(2.0), Inches(mu_x), Inches(6.5))
ln.line.color.rgb = rgb(AMBER)
ln.line.width = Pt(3)
ln.line.dash_style = 4
lab(s, mu_x, 1.6, "true μ", size=16, color="B36B00", w=1.4)
nets = [(3.5, 1.9), (4.3, 1.7), (3.2, 1.6), (4.1, 1.5), (5.7, 1.3), (3.8, 1.8)]
for i, (m, half) in enumerate(nets):
    y = 2.55 + i * 0.62
    ok = (m - half) <= mu_x <= (m + half)
    col = C_A if ok else RED
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(m - half / 1.4), Inches(y), Inches(m + half / 1.4), Inches(y))
    ln.line.color.rgb = rgb(col)
    ln.line.width = Pt(5)
    circ(s, m, y, 0.2, col, a=100, line_w=1)
lab(s, 5.8, 5.5, "missed ✗", size=15, color=RED, w=1.6)
panel(s, "In plain words", [
    ("x̄  ±  t* × s/√n", 26, C_A),
    "Sample mean, plus and minus a margin of error.",
    "Example (n = 40): 65.1 ± 3.3 = 61.9 to 68.4",
    ("95% describes the method", 19, "B36B00"),
    "Repeat many times: about 95 in 100 nets catch μ.",
    "Wider net = more sure but less precise. More data = narrower net.",
])

# ================================================================ 5 p-value
s = new_slide(notes=(
    "DRAW IT: a bell curve centred at 0 = 'the world where nothing is going on (H0)'. Mark where our result landed (t = 2.2). Shade the area beyond it, on BOTH tails.\n"
    "SAY: 'This bell shows what results luck alone would produce if H0 were true. Our result landed far out. The red area is the share of luck-only results that are at least this extreme. That share is the p-value, about 2.8% here. Small area = result would be rare if nothing were going on.'\n"
    "REMIND: it is NOT the probability that H0 is true. It is how surprising the data are if H0 is true."))
title(s, "Whiteboard 5 · p-value", "The p-value is an area in the tails")
xs = np.round(np.arange(-4, 4.01, 0.2), 1)
pdf = stats.norm.pdf(xs)
tails = np.where(np.abs(xs) >= 2.2, pdf, 0.0)
curve_chart(s, 0.6, 1.7, 6.8, 5.0, xs, [tails, pdf], ["as extreme as ours", "if nothing is going on (H₀)"],
            [RED, "B9D8D3"], "test statistic t (how many SEs from ‘no effect’)", "density", label_every=5, kind="area")
text(s, 3.3, 1.85, 4.0, 0.4, "red = results at least as extreme as ours", size=14, bold=True, color=RED)
panel(s, "In plain words", [
    ("Assume H₀: nothing is going on", 19, C_A),
    "The bell = what luck alone would produce.",
    ("Our result: t = 2.2, far out", 19, RED),
    "p = the red area ≈ 0.028 (both tails).",
    "p small → luck alone rarely does this.",
    "p is NOT the chance that H₀ is true.",
])

# ================================================================ 6 error types (Venn)
s = new_slide(notes=(
    "DRAW IT: rectangle = all the studies anyone could run. Circle T = 'H0 is actually true' (nothing there). Circle R = 'we reject H0 / say significant'. They overlap.\n"
    "READ THE REGIONS: overlap T and R = we claimed an effect that is not there = TYPE I (false alarm, probability alpha). R outside T = real effect found (power). T outside R = nothing there, nothing claimed (correct). Outside both = a real effect we missed = TYPE II.\n"
    "ASK: 'Which error is worse for a new medicine? For a spam filter?'"))
title(s, "Whiteboard 6 · Two ways to be wrong", "Venn diagram of testing errors")
frame(s, 0.6, 1.8, 6.6, 4.9)
text(s, 0.8, 1.9, 4.5, 0.4, "all studies you could run", size=15, color=MUTED)
circ(s, 3.0, 4.2, 3.2, C_A, a=30)
circ(s, 4.8, 4.2, 3.2, RED, a=30)
lab(s, 2.45, 4.2, "nothing there,\nnothing\nclaimed ✓", size=14, color=DARK, w=1.4, bold=False)
lab(s, 3.9, 4.2, "TYPE I\nfalse alarm", size=17, color=RED, w=1.6)
lab(s, 5.65, 4.2, "real effect\nfound ✓", size=16, color=DARK, w=1.6)
lab(s, 2.6, 6.25, "outside both = TYPE II (real effect missed)", size=14, color="B36B00", w=4.4)
text(s, 1.35, 2.35, 2.5, 0.4, "T: H₀ is true", size=16, bold=True, color=C_A)
text(s, 4.5, 2.35, 2.6, 0.4, "R: we reject H₀", size=16, bold=True, color=RED)
panel(s, "In plain words", [
    ("Overlap = Type I", 21, RED),
    "We said “significant” but nothing was there. Chance = α (5%).",
    ("Outside both = Type II", 21, "B36B00"),
    "A real effect, but our data didn’t show it. Fix: more data.",
    "Power = the chance of finding a real effect.",
])

# ================================================================ 7 effect size
s = new_slide(notes=(
    "DRAW IT: two bell curves for two groups. d = 0.2: almost on top of each other. d = 0.5: visibly shifted. d = 0.8: clearly apart.\n"
    "SAY: 'Cohen's d is how many standard deviations apart the two group averages are. The more the curves overlap, the smaller the effect. A p-value only says whether there is a difference; d tells you how big it is.'\n"
    "OUR DATA: heavy job vs no job d = 0.73 (clearly shifted). Hostel vs home d = 0.02 (the curves sit on top of each other)."))
title(s, "Whiteboard 7 · Effect size", "How far apart are the two bells?")
xs = np.round(np.arange(-3, 4.01, 0.25), 2)
ds = [0.2, 0.5, 0.8]
for i, d in enumerate(ds):
    x = 0.6 + i * 4.1
    ch = curve_chart(s, x, 1.75, 3.95, 3.6, xs, [stats.norm.pdf(xs, 0), stats.norm.pdf(xs, d)],
                     ["group 1", "group 2"], [C_A, C_B], "SDs", "", label_every=4)
    ch.value_axis.visible = False
    overlap = 2 * stats.norm.cdf(-d / 2) * 100
    rect(s, x, 5.45, 3.95, 1.2, fill=CARD, line=MINT)
    text(s, x + 0.2, 5.5, 3.55, 0.6, f"d = {d}", size=26, bold=True, color=DARK, font=HEAD)
    text(s, x + 0.2, 6.1, 3.55, 0.5, f"overlap ≈ {overlap:.0f}%   ({['small', 'medium', 'large'][i]})",
         size=15, color=MUTED)

# ================================================================ 8 one page summary
s = new_slide(notes="One-page summary. Ask which picture each line came from.")
title(s, "Whiteboard 8 · One page", "The whole class in six lines")
rows = [
    ("Population vs sample", "Circle inside a circle. μ is the big one, x̄ is the small one."),
    ("Sampling", "Where the dots come from matters more than how many."),
    ("Sampling distribution", "Averages bunch: SE = σ / √n."),
    ("Confidence interval", "x̄ ± t*·SE: a net that catches μ about 95% of the time."),
    ("p-value", "The tail area: how surprising is our result if nothing is going on."),
    ("Effect size", "How far apart the bells are: d = (x̄₁ − x̄₂) / s."),
]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.8), Inches(12.13), Inches(4.6)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(3.6), Inches(8.53)
for i, (a_, b_) in enumerate(rows):
    for j, v in enumerate((a_, b_)):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = rgb("FFFFFF" if i % 2 == 0 else "EEF4F2")
        for rn in c.text_frame.paragraphs[0].runs:
            rn.font.size, rn.font.name = Pt(19), BODY
            rn.font.bold = j == 0
            rn.font.color.rgb = rgb(TEAL if j == 0 else INK)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("Saved", OUT, "| slides:", len(prs.slides))
