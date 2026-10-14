"""Shared python-pptx helpers (theme, cards, charts) for the class decks. Extracted from class_05 build_deck.py."""
import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ------------------------------------------------------------------ style
DARK, LIGHT, CARD = "0F3D3E", "F6F8F7", "FFFFFF"
INK, MUTED = "17302F", "5C6F6E"
TEAL, MINT, AMBER, RED = "1F7A78", "CFE5E1", "F2A541", "C8453F"
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide(dark=False, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK if dark else LIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def rect(s, x, y, w, h, fill=CARD, line=None, rounded=True, shadow=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if not shadow:
        shp.shadow.inherit = False
    return shp


def circle(s, x, y, d, fill):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    return c


def text(s, x, y, w, h, paras, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=6):
    """paras: str | list of str | list of dict(text|runs, size, bold, color, font, italic, space)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("space", space))
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", p.get("font", font))
            f.size = Pt(o.get("size", p.get("size", size)))
            f.bold = o.get("bold", p.get("bold", bold))
            f.italic = o.get("italic", p.get("italic", False))
            f.color.rgb = rgb(o.get("color", p.get("color", color)))
    return tb


def title(s, tag, head, dark=False):
    text(s, 0.6, 0.42, 12, 0.3, tag.upper(), size=12, color=AMBER if dark else TEAL, bold=True)
    text(s, 0.6, 0.72, 12.1, 0.8, head, size=34, color="FFFFFF" if dark else DARK, font=HEAD, bold=True)


def card(s, x, y, w, h, head, body, head_color=TEAL, size=15, fill=CARD, line=MINT):
    rect(s, x, y, w, h, fill=fill, line=line)
    text(s, x + 0.25, y + 0.2, w - 0.5, 0.4, head, size=17, color=head_color, bold=True, font=HEAD)
    text(s, x + 0.25, y + 0.68, w - 0.5, h - 0.85, body, size=size, color=INK)


def style_chart(ch, cat_title, val_title):
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = BODY
    ch.font.size = Pt(12)
    ch.font.color.rgb = rgb(MUTED)
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = rgb(MINT)
    va.format.line.fill.background()
    ca.format.line.color.rgb = rgb(MUTED)
    for ax, t in ((ca, cat_title), (va, val_title)):
        ax.has_title = True
        ax.axis_title.text_frame.text = t
        r = ax.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(12), False, BODY
        r.font.color.rgb = rgb(MUTED)


