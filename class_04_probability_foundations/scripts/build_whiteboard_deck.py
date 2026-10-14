"""Builds the Class 04 whiteboard deck (Venn diagrams + maths): ../slides/Class04_Whiteboard_Venn.pptx

Six slides you can show on screen or redraw on the whiteboard. All numbers are the same as in the main deck.
"""
import os

from pptx.oxml.ns import qn
from lxml import etree

from deck_helpers import *  # noqa: F401,F403

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "slides", "Class04_Whiteboard_Venn.pptx")
C_A, C_B = "1F7A78", "E08A00"


def alpha(shape, pct):
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct * 1000)))


def venn_circle(s, cx, cy, d, color, a=35):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(color)
    alpha(c, a)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(4)
    c.shadow.inherit = False
    return c


def frame(s, x, y, w, h):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background()
    r.line.color.rgb = rgb(DARK)
    r.line.width = Pt(3)
    r.shadow.inherit = False
    return r


def label(s, x, y, t, size=26, color=DARK, w=1.6, bold=True):
    text(s, x - w / 2, y - 0.3, w, 0.6, t, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space=0)


def math_panel(s, head, lines, y=1.8, h=4.9):
    rect(s, 7.6, y, 5.13, h, fill=CARD, line=MINT)
    text(s, 7.9, y + 0.2, 4.6, 0.4, head, size=17, bold=True, color=TEAL, font=HEAD)
    paras = []
    for ln in lines:
        if isinstance(ln, tuple):
            paras.append({"text": ln[0], "size": ln[1], "bold": True, "color": ln[2] if len(ln) > 2 else DARK, "font": HEAD})
        else:
            paras.append({"text": ln, "size": 17})
    text(s, 7.9, y + 0.8, 4.6, h - 1.0, paras, space=10)


def draw_note(s, extra):
    s.notes_slide.notes_text_frame.text = extra


# ---------------------------------------------------------------- 1 sample space
s = new_slide(notes=(
    "DRAW IT: 1) a big rectangle and write 'all 36 outcomes (sample space)' in a corner. "
    "2) a circle inside it, label it A = 'first die is 6'. 3) write 6 inside the circle and 30 outside.\n"
    "SAY: 'The rectangle is everything that can happen. The circle is the event we care about. Probability is how much of the rectangle the circle covers: part divided by whole.'\n"
    "ASK: 'What is the probability of NOT A?' (30/36 = 1 - 1/6.)"))
title(s, "Whiteboard 1 · Sample space", "The rectangle is everything that can happen")
frame(s, 0.6, 1.8, 6.6, 4.9)
text(s, 0.8, 1.9, 5.5, 0.4, "all 36 outcomes of two dice (sample space)", size=15, color=MUTED)
venn_circle(s, 3.9, 4.25, 2.6, C_A)
label(s, 3.9, 3.9, "A", size=30, color=C_A)
label(s, 3.9, 4.55, "6", size=28)
text(s, 3.0, 5.7, 1.8, 0.4, "first die is 6", size=14, color=C_A, align=PP_ALIGN.CENTER)
label(s, 1.6, 6.2, "30", size=24, color=MUTED)
label(s, 6.4, 2.6, "not A", size=18, color=MUTED, w=1.2)
math_panel(s, "The maths", [
    ("P(A) = part ÷ whole", 22),
    "= 6 / 36 = 1/6",
    ("P(not A) = 1 − P(A)", 22),
    "= 1 − 1/6 = 30/36",
    "All outcomes together always add up to 1.",
])

# ---------------------------------------------------------------- 2 addition rule
s = new_slide(notes=(
    "DRAW IT: rectangle, two overlapping circles A and B. Write 3 in the left part, 3 in the middle (overlap), 3 in the right part, and 27 outside.\n"
    "SAY: 'A is first die is 6: 6 outcomes. B is total at least 10: 6 outcomes. But 3 outcomes are in both, (6,4), (6,5), (6,6). If I just add 6 + 6 I count the middle twice, so I subtract it once.'\n"
    "CHECK: 3 + 3 + 3 = 9 outcomes inside either circle; 9/36 = 0.25."))
title(s, "Whiteboard 2 · Addition rule", "A or B: don’t count the overlap twice")
frame(s, 0.6, 1.8, 6.6, 4.9)
venn_circle(s, 2.9, 4.2, 3.0, C_A)
venn_circle(s, 4.9, 4.2, 3.0, C_B)
label(s, 2.2, 4.2, "3", size=30)
label(s, 3.9, 4.2, "3", size=30)
label(s, 5.6, 4.2, "3", size=30)
label(s, 1.5, 6.25, "27 outside", size=18, color=MUTED, w=2.2)
text(s, 1.2, 2.0, 2.6, 0.5, "A: first die 6", size=17, bold=True, color=C_A)
text(s, 4.6, 2.0, 2.6, 0.5, "B: total ≥ 10", size=17, bold=True, color=C_B)
math_panel(s, "The maths", [
    ("P(A or B) = P(A) + P(B) − P(A and B)", 20),
    "P(A) = 6/36      P(B) = 6/36",
    "P(A and B) = 3/36  (the overlap)",
    ("= 6/36 + 6/36 − 3/36 = 9/36 = 0.25", 20, TEAL),
    "Check: 3 + 3 + 3 = 9 outcomes inside either circle.",
])

# ---------------------------------------------------------------- 3 conditional
s = new_slide(notes=(
    "DRAW IT: rectangle (400 students), circle G = study group (150), circle P = passed (270). Overlap 120. G-only 30. P-only 150. Outside 100.\n"
    "SAY: 'Given means: ignore everything outside the circle you are told about. If I know you are in the study group, the world is just circle G, 150 students, and 120 of them passed: 120/150.'\n"
    "THEN: 'Flip it. Given you passed, the world is circle P, 270 students, and 120 are in a group: 120/270. Same overlap on top, different circle on the bottom.'\n"
    "That is why P(A|B) is not P(B|A)."))
title(s, "Whiteboard 3 · Conditional probability", "‘Given’ means: shrink the world to one circle")
frame(s, 0.6, 1.8, 6.6, 4.9)
venn_circle(s, 3.1, 4.25, 2.9, C_A)
venn_circle(s, 4.9, 4.25, 3.4, C_B)
label(s, 2.1, 4.25, "30", size=28)
label(s, 3.9, 4.25, "120", size=30, w=1.4)
label(s, 5.75, 4.25, "150", size=28)
label(s, 1.4, 6.3, "100 outside", size=18, color=MUTED, w=2.2)
text(s, 1.2, 1.95, 2.8, 0.5, "G: study group (150)", size=16, bold=True, color=C_A)
text(s, 4.4, 1.95, 2.7, 0.5, "P: passed (270)", size=16, bold=True, color=C_B)
math_panel(s, "The maths", [
    ("P(A | B) = P(A and B) / P(B)", 22),
    "Given group: world = circle G",
    ("P(pass | group) = 120/150 = 0.80", 20, TEAL),
    "Given pass: world = circle P",
    ("P(group | pass) = 120/270 = 0.44", 20, RED),
    "Same top (120), different bottom.",
])

# ---------------------------------------------------------------- 4 independence
s = new_slide(notes=(
    "DRAW IT: a square. Split it down the middle: left half is A (first die even, 1/2). Draw a thin strip along the bottom, 1/6 of the height: that is B (total is 7, 1/6). Where they cross is A and B.\n"
    "SAY: 'If A and B are independent, B takes up the same fraction of A as it does of the whole square. The bottom strip is 1/6 of the whole square and 1/6 of the left half too. So the crossing area is 1/2 x 1/6 = 1/12. That is exactly what we computed: 3/36.'\n"
    "CONTRAST: 'Exclusive events are two circles that don't touch. P(A and B) = 0. That is not P(A) x P(B), so exclusive events are dependent.'"))
title(s, "Whiteboard 4 · Independence", "Overlap = width × height")
frame(s, 0.6, 1.8, 6.6, 4.9)
a_rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.8), Inches(3.3), Inches(4.9))
a_rect.fill.solid(); a_rect.fill.fore_color.rgb = rgb(C_A); alpha(a_rect, 30)
a_rect.line.fill.background(); a_rect.shadow.inherit = False
b_rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(5.88), Inches(6.6), Inches(0.82))
b_rect.fill.solid(); b_rect.fill.fore_color.rgb = rgb(C_B); alpha(b_rect, 40)
b_rect.line.fill.background(); b_rect.shadow.inherit = False
label(s, 2.25, 3.4, "A: first die even", size=20, color=C_A, w=3.0)
label(s, 2.25, 3.95, "width 1/2", size=18, color=C_A, w=3.0)
label(s, 5.5, 4.2, "B: total is 7", size=20, color=C_B, w=2.6)
label(s, 5.5, 4.75, "height 1/6 (bottom strip)", size=16, color=C_B, w=3.0, bold=False)
label(s, 2.25, 6.3, "A and B = 1/12", size=19, color=DARK, w=3.0)
math_panel(s, "The maths", [
    ("Independent ⇔ P(A and B) = P(A)·P(B)", 18),
    "P(A) = 1/2      P(B) = 1/6",
    "P(A) × P(B) = 1/12",
    ("P(A and B) = 3/36 = 1/12  ✅ equal", 20, TEAL),
    "Exclusive events (circles that don’t touch): P(A and B) = 0, which is NOT P(A)·P(B). So they are dependent.",
])

# ---------------------------------------------------------------- 5 Bayes
s = new_slide(notes=(
    "DRAW IT: a rectangle for 10,000 people. A small circle S = sick (100 people). A big circle + = tests positive (590 people), overlapping S almost completely. Write 95 in the overlap, 5 in the sliver of S outside +, 495 in the part of + outside S, 9,405 outside both. (Not to scale: say so.)\n"
    "SAY: 'You tested positive, so you are somewhere inside the + circle. 590 people are in that circle. How many of them are sick? 95. So 95 out of 590, 16%. The 495 are healthy people the test falsely alarmed. The healthy population is huge, so even a 5% false-alarm rate fills the + circle with healthy people.'\n"
    "KEY: the denominator is the WHOLE positive circle, not the sick circle."))
title(s, "Whiteboard 5 · Bayes", "You tested positive: which circle are you in?")
frame(s, 0.6, 1.8, 6.6, 4.9)
text(s, 0.8, 1.9, 5.5, 0.4, "10,000 people (not to scale)", size=15, color=MUTED)
venn_circle(s, 3.65, 4.3, 1.9, C_A, a=40)
venn_circle(s, 5.0, 4.25, 3.6, C_B, a=30)
label(s, 2.95, 4.3, "5", size=20, w=0.8)
label(s, 4.0, 4.3, "95", size=26, w=1.0)
label(s, 5.9, 4.25, "495", size=26, w=1.2)
label(s, 1.5, 6.3, "9,405 outside", size=18, color=MUTED, w=2.4)
text(s, 0.9, 2.35, 2.6, 0.5, "S: sick (100)", size=17, bold=True, color=C_A)
text(s, 4.2, 2.0, 2.9, 0.5, "+ : tests positive (590)", size=17, bold=True, color=C_B)
math_panel(s, "The maths", [
    ("P(sick | +) = P(+ | sick) · P(sick) / P(+)", 19),
    "Count version: inside the + circle,",
    ("95 sick out of 95 + 495 = 590", 20, TEAL),
    ("= 95 / 590 = 0.161", 22, RED),
    "Formula version:",
    "0.95 × 0.01 = 0.0095 (true +)",
    "0.05 × 0.99 = 0.0495 (false +)",
    "0.0095 / (0.0095 + 0.0495) = 0.161",
])

# ---------------------------------------------------------------- 6 formula page
s = new_slide(notes=("Use as a one-page summary at the end, or leave it on screen while students work on the tasks. "
                     "Ask them to point at which Venn picture each formula came from."))
title(s, "Whiteboard 6 · All the formulas", "Everything on one page")
rows = [
    ("Complement", "P(not A) = 1 − P(A)"),
    ("Addition", "P(A or B) = P(A) + P(B) − P(A and B)"),
    ("Conditional", "P(A | B) = P(A and B) / P(B)"),
    ("Multiplication", "P(A and B) = P(A | B) · P(B)"),
    ("Independence", "P(A and B) = P(A) · P(B)"),
    ("Total probability", "P(E) = P(E | H) P(H) + P(E | not H) P(not H)"),
    ("Bayes", "P(H | E) = P(E | H) · P(H) / P(E)"),
    ("Expected value", "E[X] = Σ x · P(x)"),
    ("Variance", "Var(X) = Σ (x − E[X])² · P(x)"),
]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.7), Inches(12.13), Inches(5.0)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(3.4), Inches(8.73)
for i, (a, b) in enumerate(rows):
    for j, v in enumerate((a, b)):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = rgb("FFFFFF" if i % 2 == 0 else "EEF4F2")
        for rn in c.text_frame.paragraphs[0].runs:
            rn.font.size, rn.font.name = Pt(19), (HEAD if j else BODY)
            rn.font.bold = j == 0
            rn.font.color.rgb = rgb(TEAL if j == 0 else INK)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("Saved", OUT, "| slides:", len(prs.slides))
